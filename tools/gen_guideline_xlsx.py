#!/usr/bin/env python3
"""正本md（UI/UXレビューガイドライン）から閲覧・記入用xlsxを生成する。

使い方:
    python3 gen_guideline_xlsx.py uiux_review_guideline_rev1.md
    -> uiux_review_guideline_rev1.xlsx を同じ場所に出力
必要: pip install openpyxl pyyaml
"""
import re
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Meiryo UI"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
THIN = Side(style="thin", color="C8CDD3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def load(md_path: Path):
    text = md_path.read_text(encoding="utf-8")
    front = yaml.safe_load(text.split("---", 2)[1])
    items = None
    for block in re.findall(r"```yaml\n(.*?)```", text, re.S):
        data = yaml.safe_load(block)
        if isinstance(data, dict) and "items" in data:
            items = data["items"]
            break
    if items is None:
        sys.exit("items の yaml ブロックが見つかりません")
    principles = [
        [c.strip() for c in line.strip().strip("|").split("|")]
        for line in text.splitlines()
        if re.match(r"^\|\s*P\d{2}\s*\|", line)
    ]
    return front, items, principles


def header(ws, cols):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 30


def body_cell(ws, r, c, v, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, size=10)
    cell.alignment = WRAP
    cell.border = BORDER
    if fill:
        cell.fill = fill
    return cell


def main():
    md = Path(sys.argv[1] if len(sys.argv) > 1 else "uiux_review_guideline_rev1.md")
    front, items, principles = load(md)
    wb = Workbook()

    # --- 読み方 ---
    ws = wb.active
    ws.title = "読み方"
    lines = [
        f"{front['doc']}（{front['version']}）",
        f"このファイルは {md.name} から自動生成された閲覧・記入用です。項目の追加・修正は正本md側で行い、再生成してください。",
        "",
        "【記入方法】チェックリストシートの黄色の列（判定・深刻度・画面・メモ）に記入します。",
        "・判定: OK / NG / 対象外",
        "・深刻度: NGのとき 0〜4 を入力（空欄なら既定値を採用）。対応優先度は自動で表示されます",
        "・method: auto=ツール判定 / llm=LLM一次判定→人が確認 / human=人のみ",
        "",
        f"主な拠り所: {front['primary_reference']}",
    ]
    for i, t in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, size=12 if i == 1 else 10, bold=(i == 1))
    ws.column_dimensions["A"].width = 120

    # --- チェックリスト ---
    ws = wb.create_sheet("チェックリスト")
    cols = [
        ("ID", 9), ("カテゴリ", 8), ("タイトル", 26), ("チェック内容", 50),
        ("OK例", 34), ("NG例", 34), ("既定深刻度", 8), ("フェーズ", 11),
        ("method", 8), ("ツール", 18), ("原則", 9), ("根拠", 30),
        ("判定", 9), ("深刻度", 8), ("採用深刻度", 9), ("対応優先度", 22),
        ("画面", 18), ("メモ", 30),
    ]
    header(ws, cols)
    for r, it in enumerate(items, 2):
        vals = [
            it["id"], it["category"], it["title"], it["check"], it.get("ok", ""),
            it.get("ng", ""), it["severity_default"], ", ".join(it["phase"]),
            it["method"], it.get("tool", ""), ", ".join(it.get("principle", [])),
            "\n".join(it.get("refs", [])),
        ]
        for c, v in enumerate(vals, 1):
            body_cell(ws, r, c, v)
        for c in (13, 14, 17, 18):
            body_cell(ws, r, c, None, INPUT_FILL)
        body_cell(ws, r, 15, f'=IF(M{r}="NG",IF(N{r}="",G{r},N{r}),"")')
        body_cell(
            ws, r, 16,
            f'=IF(O{r}="","",CHOOSE(O{r}+1,"対応不要","バックログ","次回以降に計画修正",'
            f'"リリース前に修正","リリース不可・即時修正"))',
        )
    last = len(items) + 1
    dv1 = DataValidation(type="list", formula1='"OK,NG,対象外"', allow_blank=True)
    dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"M2:M{last}")
    dv2.add(f"N2:N{last}")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # --- 原則 ---
    ws = wb.create_sheet("L1原則")
    header(ws, [("ID", 8), ("原則", 22), ("要旨", 70), ("主な出典", 36)])
    for r, row in enumerate(principles, 2):
        for c, v in enumerate(row, 1):
            body_cell(ws, r, c, v)

    # --- 深刻度 ---
    ws = wb.create_sheet("深刻度")
    header(ws, [("深刻度", 8), ("定義", 60), ("対応", 40)])
    sev = [
        (4, "主要タスクが完了できない／データ損失・誤処理・情報漏えい", "リリース不可・即時修正"),
        (3, "主要タスクの完了に大きな支障。多くのユーザーが詰まる", "リリース前に修正"),
        (2, "手間・迷いが生じるが完了できる", "次回以降に計画修正"),
        (1, "見た目・表現の問題。業務影響ほぼなし", "バックログ"),
        (0, "ユーザビリティ上の問題ではない", "対応不要"),
    ]
    for r, row in enumerate(sev, 2):
        for c, v in enumerate(row, 1):
            body_cell(ws, r, c, v)

    out = md.with_suffix(".xlsx")
    wb.save(out)
    print(f"{out} を出力しました（{len(items)}項目）")


if __name__ == "__main__":
    main()
